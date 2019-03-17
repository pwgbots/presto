# Software developed by Pieter W.G. Bots for the PrESTO project
# Code repository: https://github.com/pwgbots/presto
# Project wiki: http://presto.tudelft.nl/wiki

"""
Copyright (c) 2019 Delft University of Technology

Permission is hereby granted, free of charge, to any person obtaining a copy of
this software and associated documentation files (the "Software"), to deal in
the Software without restriction, including without limitation the rights to use,
copy, modify, merge, publish, distribute, sublicense, and/or sell copies of the
Software, and to permit persons to whom the Software is furnished to do so,
subject to the following conditions:

The above copyright notice and this permission notice shall be included in all
copies or substantial portions of the Software.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR IMPLIED,
INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY, FITNESS FOR A
PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE AUTHORS OR COPYRIGHT
HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER LIABILITY, WHETHER IN AN ACTION
OF CONTRACT, TORT OR OTHERWISE, ARISING FROM, OUT OF OR IN CONNECTION WITH THE
SOFTWARE OR THE USE OR OTHER DEALINGS IN THE SOFTWARE.
"""


# -*- coding: utf-8 -*-
from __future__ import unicode_literals

from django.conf import settings
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Count, Q
from django.shortcuts import render
from django.utils import timezone
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt

# python modules
from docx import Document
from openpyxl import load_workbook
import os
from pptx import Presentation
from PyPDF2 import PdfFileReader
from tempfile import mkstemp

# presto modules
from presto.generic import generic_context, warn_user, inform_user

MAX_UPLOAD_SIZE = 2.0  # max 2 MB per upload file


# view to test whether upload file is valid
@method_decorator(csrf_exempt, name='dispatch')
@login_required(login_url=settings.LOGIN_URL)
def test_file_type(request, **kwargs):
    context = generic_context(request)
    context['page_title'] = 'Presto'
    if not request.FILES:
        return render(request, 'presto/test_file_type.html', context)
    # check whether uploaded file is valid
    f = request.FILES['test']
    upload_size = f.size / 1048576.0 # size in MB 
    ext = os.path.splitext(f.name)[1]
    inform_user(context, 'File uploaded', 'Size: %3.2f MB<br/>Extension: %s' % (upload_size, ext))
    file_too_big = upload_size > MAX_UPLOAD_SIZE
    if file_too_big:
      warn_user(context, 'File too big', 'Size exceeds %3.2f MB' % MAX_UPLOAD_SIZE)
    else:
        try:
            # save upload as temporary file
            handle, fn = mkstemp()
            os.close(handle)
            with open(fn, 'wb+') as dst:
                for chunk in f.chunks():
                    dst.write(chunk)
            # NOTE: openpyxl requires files to have a valid extension!
            # NOTE: extension is always added for uniformity & tracing purposes
            os.rename(fn, fn + ext)
            fn += ext
            # try to open the file with the correct application
            if ext in ['.docx']:
                doc = Document(fn)
            elif ext in ['.pptx']:
                prs = Presentation(fn)
            elif ext in ['.xlsx']:
                wbk = load_workbook(fn)
            elif ext == '.pdf':
                pdf = PdfFileReader(fn)
        # NOTE: assumes that incorrect file types will raise an error
        except Exception, e:
            warn_user(context, 'Invalid_file', 'ERROR while checking file %s: %s' % (fn, str(e)))

    return render(request, 'presto/test_file_type.html', context)
